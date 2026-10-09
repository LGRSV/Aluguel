"""
Pós-coleta do aluguel: classifica com o Haiku, gera planilha e uma página única (fotos embutidas).

    dados/classificacao.json            o que o Haiku extraiu de cada anúncio
    planilhas/aluguel_palmas_<data>.xlsx  anúncios + médias por tipo/quartos e por região
    planilhas/historico_aluguel.xlsx    uma linha por coleta (média geral e por tipo)
    index.html                          página com filtros e fotos (publicada no GitHub)
    historico/anuncios.json             quando cada anúncio apareceu e saiu do ar (ver ../marketplace-iphone13/historico.py)

Uso:  python gerar_aluguel.py            (a coleta_aluguel.py já chama no fim)
      python gerar_aluguel.py --sem-haiku  (reaproveita dados/classificacao.json)
"""

import base64
import glob
import io
import json
import os
import re
import statistics
import subprocess
import sys
from collections import defaultdict
from datetime import date, timedelta

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from PIL import Image

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(os.path.dirname(AQUI), "marketplace-iphone13"))
import historico  # noqa: E402
DADOS = os.path.join(AQUI, "dados")
PLAN = os.path.join(AQUI, "planilhas")
LOTE = 30
REGIAO_OK = ("palmas", "taquaralto", "taquaruçu", "taquarucu", "porto nacional", "luzimangues")

PROMPT = """Você classifica anúncios de imóveis do Facebook Marketplace de Palmas/TO. Responda SÓ com um array JSON,
um objeto por anúncio, na mesma ordem, sem texto antes ou depois:
{"id":"...","finalidade":"aluguel"|"temporada"|"venda"|"outro",
 "tipo":"casa"|"apartamento"|"kitnet"|"quarto"|"sobrado"|"sala comercial"|"galpão"|"terreno"|"outro",
 "quartos":número ou null,"banheiros":número ou null,"area_m2":número ou null,
 "mobiliado":"sim"|"semi"|"não"|null,"condominio_incluso":true|false|null,
 "bairro":"quadra ou bairro como escrito (ex.: 606 Sul, ARSE 51, Taquaralto, Aureny III), ou vazio",
 "regiao":"Plano Diretor Sul"|"Plano Diretor Norte"|"Região Sul (Taquaralto/Aurenys)"|"Taquaruçu"|"Luzimangues"|"Porto Nacional"|"outra"|"",
 "preco_mensal":número ou null,"observacoes":"até 120 caracteres"}
Regras: "aluguel" é locação mensal residencial ou comercial; diária, evento ou chácara por fim de semana é
"temporada"; parcelado, entrada, ágio ou venda é "venda". preco_mensal é o aluguel por mês (null se não for mensal).
Quadras de Palmas: ARSE/ARSO/números terminados em Sul = Plano Diretor Sul; ARNE/ARNO/Norte = Plano Diretor Norte.
Nunca invente quartos, área ou bairro: deixe null/vazio se não estiver escrito.
Anúncios:
"""


def carregar():
    itens, vistos = [], set()
    for arq in sorted(glob.glob(os.path.join(DADOS, "_anuncios_*.json"))):
        for it in json.load(open(arq, encoding="utf-8-sig")):
            if it.get("erro") or it["id"] in vistos:
                continue
            vistos.add(it["id"])
            it["preco"] = int(re.sub(r"\D", "", str(it.get("preco") or "0")) or 0)
            itens.append(it)
    return itens


def haiku(lote):
    entrada = PROMPT + json.dumps([{k: i.get(k, "") for k in ("id", "titulo", "preco", "local", "detalhes")}
                                   | {"descricao": i.get("descricao", "")[:600]} for i in lote], ensure_ascii=False)
    r = subprocess.run(["claude", "-p", "--model", "haiku"], input=entrada, capture_output=True, text=True,
                       encoding="utf-8", timeout=400, shell=(os.name == "nt"))
    m = re.search(r"\[.*\]", r.stdout, re.S)
    if not m:
        raise RuntimeError(f"sem JSON: {r.stdout[:200]} {r.stderr[:200]}")
    return {str(c["id"]): c for c in json.loads(m.group(0)) if c.get("id")}


def classificar(itens):
    cls = {}
    for k in range(0, len(itens), LOTE):
        lote = itens[k:k + LOTE]
        try:
            cls.update(haiku(lote))
        except Exception as e:
            print(f"  lote {k // LOTE + 1} falhou ({e}); de novo")
            cls.update(haiku(lote))
        print(f"  lote {k // LOTE + 1}/{-(-len(itens) // LOTE)} ok")
    json.dump(cls, open(os.path.join(DADOS, "classificacao.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    return cls


def media(v):
    return round(statistics.mean(v)) if v else None


def faixa_quartos(q):
    return "—" if not q else ("1 quarto" if q == 1 else (f"{q} quartos" if q < 4 else "4+ quartos"))


def planilha(validos, todos, hoje):
    os.makedirs(PLAN, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    ws["A1"] = f"Aluguel em Palmas/TO — Facebook Marketplace — coleta de {hoje:%d/%m/%Y}"
    ws["A1"].font = Font(size=14, bold=True, color="1F3864")
    ws["A2"] = f"{len(validos)} anúncios de aluguel mensal na região (de {len(todos)} lidos). Médias de preço anunciado."
    linha = 4
    for titulo, chave in (("Por tipo e quartos", lambda i: (i["tipo"], faixa_quartos(i["quartos"]))),
                          ("Por região", lambda i: (i["regiao"] or "não informado", "")),
                          ("Por tipo", lambda i: (i["tipo"], ""))):
        ws.cell(linha, 1, titulo).font = Font(bold=True, size=12)
        linha += 1
        for c, h in enumerate(("Grupo", "", "Qtd", "Média", "Menor", "Maior"), 1):
            ws.cell(linha, c, h).font = Font(bold=True)
        linha += 1
        g = defaultdict(list)
        for i in validos:
            g[chave(i)].append(i["mensal"])
        for (a, b), v in sorted(g.items(), key=lambda x: (-len(x[1]), x[0])):
            ws.append([a, b, len(v), media(v), min(v), max(v)])
            if len(v) < 5:
                ws.cell(ws.max_row, 3).fill = PatternFill("solid", fgColor="FFF2CC")
            linha = ws.max_row + 1
        linha += 1
    for col, w in zip("ABCDEF", (30, 14, 8, 12, 12, 12)):
        ws.column_dimensions[col].width = w
    a = wb.create_sheet("Anúncios")
    cab = ["Tipo", "Quartos", "Banheiros", "Área m²", "Mobiliado", "Bairro", "Região", "Aluguel/mês", "Título", "Local", "Finalidade", "Observações", "Link"]
    a.append(cab)
    for c in range(1, len(cab) + 1):
        a.cell(1, c).font = Font(bold=True)
    for i in sorted(todos, key=lambda x: (x["finalidade"] != "aluguel", x["mensal"] or 0)):
        a.append([i["tipo"], i["quartos"], i["banheiros"], i["area_m2"], i["mobiliado"], i["bairro"], i["regiao"],
                  i["mensal"], i["titulo"], i["local"], i["finalidade"], i["obs"], i["link"]])
        if i["finalidade"] != "aluguel" or not i["na_regiao"]:
            for c in range(1, len(cab) + 1):
                a.cell(a.max_row, c).font = Font(color="999999")
    arq = os.path.join(PLAN, f"aluguel_palmas_{hoje.isoformat()}.xlsx")
    wb.save(arq)

    hist = os.path.join(PLAN, "historico_aluguel.xlsx")
    tipos = ["casa", "apartamento", "kitnet", "sobrado", "quarto", "sala comercial"]
    if os.path.exists(hist):
        hw = load_workbook(hist)
        hs = hw.active
    else:
        hw = Workbook()
        hs = hw.active
        hs.title = "Historico"
        hs.append(["Data", "Anúncios", "Média geral"] + [f"Média {t}" for t in tipos])
    hs.append([hoje.strftime("%d/%m/%Y"), len(validos), media([i["mensal"] for i in validos])]
              + [media([i["mensal"] for i in validos if i["tipo"] == t]) for t in tipos])
    hw.save(hist)
    return arq


def tira(fotos_b64, lado=380, maxf=5):
    fotos = fotos_b64[:maxf]
    if not fotos:
        return "", 0
    t = Image.new("RGB", (lado * len(fotos), lado), (0, 0, 0))
    for k, b in enumerate(fotos):
        im = Image.open(io.BytesIO(base64.b64decode(b))).convert("RGB")
        im.thumbnail((lado, lado))
        t.paste(im, (k * lado + (lado - im.width) // 2, (lado - im.height) // 2))
    buf = io.BytesIO()
    t.save(buf, "JPEG", quality=52, optimize=True)
    return "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode(), len(fotos)


def pagina(validos, hoje):
    grupos = defaultdict(list)
    for i in validos:
        grupos[(i["tipo"], faixa_quartos(i["quartos"]))].append(i["mensal"])
    D = []
    for i in validos:
        v = grupos[(i["tipo"], faixa_quartos(i["quartos"]))]
        vs = round((i["mensal"] / media(v) - 1) * 100) if len(v) >= 4 else None
        src, nf = tira(i["fotos"])
        D.append({"id": i["id"], "tipo": i["tipo"], "q": faixa_quartos(i["quartos"]), "regiao": i["regiao"] or "não informado",
                  "bairro": i["bairro"], "preco": i["mensal"], "titulo": i["titulo"], "mob": i["mobiliado"] or "",
                  "area": i["area_m2"], "desc": i["descricao"], "obs": i["obs"], "link": i["link"], "vs": vs,
                  "media": media(v), "n": len(v), "tira": src, "nf": nf, "novo": i.get("novo", False)})
    medias = sorted(([f"{t} · {q}", media(v), len(v)] for (t, q), v in grupos.items() if len(v) >= 2), key=lambda x: -x[2])
    historico.definir_modelos(AQUI, {i["id"]: (i["tipo"], f"{i['tipo']} · {faixa_quartos(i['quartos'])}") for i in validos})
    resumo = {"medias": medias, "vendas": historico.resumo(historico.carregar(AQUI))}
    html = MODELO.replace("__DADOS__", json.dumps(D, ensure_ascii=False).replace("</", "<\\/")) \
        .replace("__RESUMO__", json.dumps(resumo, ensure_ascii=False)).replace("__DATA__", f"{hoje:%d/%m/%Y}")
    arq = os.path.join(AQUI, "index.html")
    open(arq, "w", encoding="utf-8").write(html)
    return arq, len(html.encode()) / 1e6


def main():
    itens = carregar()
    print(f"{len(itens)} anúncios lidos")
    caminho = os.path.join(DADOS, "classificacao.json")
    cls = json.load(open(caminho, encoding="utf-8")) if os.path.exists(caminho) else {}
    faltam = [i for i in itens if i["id"] not in cls]
    if faltam and "--sem-haiku" not in sys.argv:  # incremental: só os anúncios ainda não classificados
        print(f"{len(faltam)} anúncios novos para o Haiku")
        cls = {**cls, **classificar(faltam)}
        json.dump(cls, open(caminho, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    todos = []
    for i in itens:
        c = cls.get(i["id"], {})
        cidade = (i.get("local") or "").split(",")[-2].strip().lower() if "," in (i.get("local") or "") else ""
        todos.append({**i, "finalidade": c.get("finalidade") or "outro", "tipo": c.get("tipo") or "outro",
                      "quartos": c.get("quartos"), "banheiros": c.get("banheiros"), "area_m2": c.get("area_m2"),
                      "mobiliado": c.get("mobiliado"), "bairro": c.get("bairro") or "", "regiao": c.get("regiao") or "",
                      "mensal": c.get("preco_mensal") or (i["preco"] if c.get("finalidade") == "aluguel" else None),
                      "obs": c.get("observacoes", ""), "na_regiao": cidade in REGIAO_OK,
                      "link": f"https://www.facebook.com/marketplace/item/{i['id']}/"})
    validos = [i for i in todos if i["finalidade"] == "aluguel" and i["na_regiao"] and i["mensal"] and 150 <= i["mensal"] <= 50000]
    hoje = date.today()
    H = historico.carregar(AQUI)
    validos = [i for i in validos if H.get(i["id"], {}).get("status", "ativo") == "ativo"]  # alugados/apagados saem da página
    if "--diario" not in sys.argv:  # planilhas só na coleta completa (o histórico do Excel é semanal)
        print("planilha:", planilha(validos, todos, hoje))
    for i in validos:
        i["novo"] = H.get(i["id"], {}).get("primeiro_visto", "") >= (hoje - timedelta(days=1)).isoformat()
    arq, mb = pagina(validos, hoje)
    print(f"página: {arq} ({mb:.1f} MB)")
    por_tipo = defaultdict(list)
    for i in validos:
        por_tipo[i["tipo"]].append(i["mensal"])
    print(f"{len(validos)} aluguéis válidos na região:")
    for t, v in sorted(por_tipo.items(), key=lambda x: -len(x[1])):
        print(f"  {t:16} {len(v):>4}  média R$ {media(v):>6}")


MODELO = r"""<!doctype html>
<html lang="pt-BR"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Aluguel em Palmas</title>
<style>
:root{--bg:#f6f7f9;--card:#fff;--tx:#1c1e21;--mut:#65676b;--bd:#dadde1;--ac:#1877f2;--ok:#1a7f37;--bad:#c62828;--warn:#b26a00}
@media (prefers-color-scheme:dark){:root:not([data-theme="light"]){--bg:#18191a;--card:#242526;--tx:#e4e6eb;--mut:#b0b3b8;--bd:#3a3b3c;--ac:#4599ff;--ok:#4cc26a;--bad:#ff6b6b;--warn:#ffb74d}}
:root[data-theme="dark"]{--bg:#18191a;--card:#242526;--tx:#e4e6eb;--mut:#b0b3b8;--bd:#3a3b3c;--ac:#4599ff;--ok:#4cc26a;--bad:#ff6b6b;--warn:#ffb74d}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--tx);font:15px/1.4 system-ui,Segoe UI,Roboto,sans-serif}
header{position:sticky;top:0;z-index:5;background:var(--card);border-bottom:1px solid var(--bd);padding:12px 16px}
h1{margin:0 0 4px;font-size:20px}.sub{color:var(--mut);font-size:13px;margin-bottom:10px}
.bar{display:flex;flex-wrap:wrap;gap:8px}.bar select,.bar input{font:inherit;padding:6px 12px;border-radius:18px;border:1px solid var(--bd);background:var(--bg);color:var(--tx);max-width:100%}
.bar input{flex:1;min-width:150px}
.wrap{padding:0 16px;max-width:1400px;margin:0 auto}
details{margin-top:12px}summary{cursor:pointer;font-weight:600}
table{border-collapse:collapse;font-size:14px;margin:8px 0}td,th{border:1px solid var(--bd);padding:4px 10px;text-align:right}td:first-child,th:first-child{text-align:left}
main{display:grid;grid-template-columns:repeat(auto-fill,minmax(220px,1fr));gap:14px;padding:12px 0}
.c{background:var(--card);border:1px solid var(--bd);border-radius:10px;overflow:hidden;cursor:pointer}
.ph{width:100%;aspect-ratio:1;background-color:#000;background-repeat:no-repeat}
.b{padding:10px}.p{font-weight:700;font-size:18px}.t{margin:2px 0 4px;display:-webkit-box;-webkit-line-clamp:2;-webkit-box-orient:vertical;overflow:hidden}
.m{color:var(--mut);font-size:13px}.tag{font-size:12px;padding:2px 8px;border-radius:10px;border:1px solid var(--bd);display:inline-block;margin-top:6px}
.low{color:var(--ok);border-color:var(--ok)}.high{color:var(--bad);border-color:var(--bad)}
.n{position:relative}.n span{position:absolute;right:8px;bottom:8px;background:#000a;color:#fff;font-size:12px;padding:2px 7px;border-radius:10px}
#mod{position:fixed;inset:0;background:#000c;display:none;z-index:10;padding:16px;overflow:auto}
#mod .box{background:var(--card);max-width:1000px;margin:0 auto;border-radius:12px;display:grid;grid-template-columns:1.2fr 1fr;overflow:hidden}
@media (max-width:760px){#mod .box{grid-template-columns:1fr}#mod{padding:8px}}
.gal{background:#000;position:relative}.gal .ph{max-height:70vh}
.nav{position:absolute;top:40%;background:#0008;color:#fff;border:0;font-size:28px;width:44px;height:44px;border-radius:50%;cursor:pointer}
.prev{left:8px}.next{right:8px}.det{padding:18px;overflow:auto;max-height:85vh}.det h2{margin:0 0 6px;font-size:19px}
.desc{white-space:pre-wrap;margin-top:12px;border-top:1px solid var(--bd);padding-top:12px;overflow-wrap:anywhere}
.acts a{display:inline-block;margin-top:12px;background:var(--ac);color:#fff;text-decoration:none;padding:8px 14px;border-radius:8px;font-weight:600}
.x{float:right;background:none;border:0;font-size:26px;color:var(--mut);cursor:pointer}
.vazio{grid-column:1/-1;text-align:center;color:var(--mut);padding:40px}
</style></head><body>
<header><h1>Aluguel em Palmas — Facebook Marketplace</h1>
<div class="sub" id="sub"></div>
<div class="bar"><select id="ftipo"></select><select id="fq"></select><select id="freg"></select>
<select id="ord"><option value="p">Menor aluguel</option><option value="P">Maior aluguel</option><option value="v">Mais abaixo da média</option><option value="n">Novos primeiro</option></select>
<input id="q" placeholder="Buscar (bairro, mobiliado, piscina...)"></div></header>
<div class="wrap">
<details><summary>Aluguel médio por tipo e quartos</summary><div id="res"></div></details>
<details><summary>📈 O que aluga mais rápido (histórico)</summary><div id="vendas"></div></details>
<main id="g"></main></div>
<div id="mod"><div class="box"><div class="gal"><div class="ph" id="big"></div><button class="nav prev" id="pv">‹</button><button class="nav next" id="nx">›</button></div><div class="det" id="det"></div></div></div>
<script>
const D=__DADOS__,R=__RESUMO__;
const fmt=v=>'R$ '+Number(v).toLocaleString('pt-BR'),esc=s=>String(s??'').replace(/[&<>"]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}[c]));
const bg=(d,i)=>d.nf?`background-image:url(${d.tira});background-size:${d.nf*100}% 100%;background-position:${d.nf>1?i/(d.nf-1)*100:0}% 0`:'';
document.getElementById('sub').textContent=`Coleta de __DATA__ · ${D.length} imóveis para alugar (mensal) · confira no Facebook se o anúncio ainda está ativo`;
document.getElementById('res').innerHTML='<table><tr><th>Tipo · quartos</th><th>Média</th><th>Anúncios</th></tr>'+R.medias.map(([k,m,n])=>`<tr><td>${esc(k)}</td><td>${fmt(m)}</td><td>${n}${n<5?' ⚠':''}</td></tr>`).join('')+'</table><div class="m">⚠ menos de 5 anúncios: média pouco confiável.</div>';
const opts=(id,nome,campo)=>{const v=[...new Set(D.map(d=>d[campo]))].sort();document.getElementById(id).innerHTML=`<option value="">${nome}</option>`+v.map(x=>`<option>${esc(x)}</option>`).join('')};
opts('ftipo','Todos os tipos','tipo');opts('fq','Quartos','q');opts('freg','Todas as regiões','regiao');
const tag=d=>d.vs==null?'':`<span class="tag ${d.vs<0?'low':'high'}">${d.vs>0?'+':''}${d.vs}% vs média (${fmt(d.media)})</span>`;
let cur=null,idx=0;
function lista(){const t=ftipo.value,q=fq.value,r=freg.value,b=document.getElementById('q').value.toLowerCase(),o=ord.value;
 let L=D.filter(d=>(!t||d.tipo===t)&&(!q||d.q===q)&&(!r||d.regiao===r)&&(!b||(d.titulo+' '+d.desc+' '+d.bairro+' '+d.mob).toLowerCase().includes(b)));
 L.sort(o==='p'?(a,b)=>a.preco-b.preco:o==='P'?(a,b)=>b.preco-a.preco:o==='n'?(a,b)=>b.novo-a.novo||a.preco-b.preco:(a,b)=>(a.vs??99)-(b.vs??99));
 g.innerHTML=L.length?L.map(d=>`<div class="c" data-id="${d.id}"><div class="n"><div class="ph" style="${bg(d,0)}"></div><span>📷 ${d.nf}</span></div><div class="b"><div class="p">${fmt(d.preco)}<span class="m">/mês</span>${d.novo?' <span class="tag low">🆕 novo</span>':''}</div><div class="t">${esc(d.titulo)}</div><div class="m">${esc(d.tipo)} · ${esc(d.q)}${d.bairro?' · '+esc(d.bairro):''}</div>${tag(d)}</div></div>`).join(''):'<div class="vazio">Nada encontrado.</div>'}
function abrir(id){cur=D.find(d=>d.id===id);
 det.innerHTML=`<button class="x" id="fx">×</button><h2>${esc(cur.titulo)}</h2><div class="p">${fmt(cur.preco)}/mês</div>
 <div class="m">${esc(cur.tipo)} · ${esc(cur.q)}${cur.area?' · '+cur.area+' m²':''}${cur.mob?' · mobiliado: '+esc(cur.mob):''}<br>${esc(cur.bairro||'')} ${cur.regiao!=='não informado'?'('+esc(cur.regiao)+')':''}</div>${tag(cur)}
 ${cur.obs?`<div class="m" style="margin-top:6px">${esc(cur.obs)}</div>`:''}<div class="acts"><a href="${cur.link}" target="_blank" rel="noopener">Abrir no Facebook</a></div><div class="desc">${esc(cur.desc||'(sem descrição)')}</div>`;
 fx.onclick=fechar;foto(0);mod.style.display='block';document.body.style.overflow='hidden'}
function foto(i){idx=(i+Math.max(cur.nf,1))%Math.max(cur.nf,1);big.style.cssText=bg(cur,idx);pv.style.display=nx.style.display=cur.nf>1?'':'none'}
function fechar(){mod.style.display='none';document.body.style.overflow='';cur=null}
['ftipo','fq','freg','ord'].forEach(i=>document.getElementById(i).onchange=lista);document.getElementById('q').oninput=lista;
g.onclick=e=>{const c=e.target.closest('.c');if(c)abrir(c.dataset.id)};pv.onclick=()=>foto(idx-1);nx.onclick=()=>foto(idx+1);
mod.onclick=e=>{if(e.target.id==='mod')fechar()};document.onkeydown=e=>{if(!cur)return;if(e.key==='Escape')fechar();if(e.key==='ArrowLeft')foto(idx-1);if(e.key==='ArrowRight')foto(idx+1)};
{const V=R.vendas||{linhas:[]},L=V.linhas.filter(l=>l.vendidos+l.apagados>0),br=d=>d?d.split('-').reverse().join('/'):'',f=v=>v==null?'–':fmt(v);
 document.getElementById('vendas').innerHTML=!L.length?`<div class="m">Histórico começou em ${br(V.inicio)}. O ranking aparece depois da próxima coleta, quando os anúncios que saíram do ar forem conferidos.</div>`:
 '<table><tr><th>Tipo · quartos</th><th>Alugados</th><th>Apagados</th><th>Ainda ativos</th><th>Giro</th><th>Dias até sair</th><th>Aluguel de quem saiu</th><th>Aluguel dos ativos</th></tr>'+
 L.map(l=>`<tr><td>${esc(l.mod)}</td><td>${l.vendidos}</td><td>${l.apagados}</td><td>${l.ativos}</td><td>${l.giro}%</td><td>${l.dias??'–'}</td><td>${f(l.preco_saiu)}</td><td>${f(l.preco_ativo)}</td></tr>`).join('')+
 `</table><div class="m">Desde ${br(V.inicio)}. <b>Alugado</b> = marcado como indisponível no Facebook. <b>Apagado</b> = anúncio removido (pode ter alugado por fora ou desistido).</div>`}
lista();
</script></body></html>
"""

if __name__ == "__main__":
    main()
